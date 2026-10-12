"""Safe, Unicode-preserving exports of the stored transcript (no invented summary)."""

from __future__ import annotations

import math
import os
import re
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

LANGUAGES = {"auto": "Автоопределение", "kk": "Қазақша", "ru": "Русский", "en": "English"}
SOURCES = {"teams": "Microsoft Teams", "upload": "Загруженная запись"}


def _clean(value: object) -> str:
    """Remove XML-disallowed controls while retaining all normal Unicode text."""
    return "".join(
        c
        for c in str(value or "")
        if c in "\n\r\t"
        or (ord(c) >= 32 and ord(c) not in (0xFFFE, 0xFFFF) and not 0xD800 <= ord(c) <= 0xDFFF)
    )


def _timestamp(value: object) -> str:
    try:
        seconds = float(value or 0)
        seconds = max(0, int(seconds)) if math.isfinite(seconds) else 0
    except (TypeError, ValueError, OverflowError):
        seconds = 0
    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def _date(value: object) -> str:
    raw = _clean(value)
    if not raw:
        return "Не указано"
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).isoformat(
            sep=" ", timespec="seconds"
        )
    except ValueError:
        return raw


def _filename(title: str, format: str) -> str:
    # ASCII filename works with ordinary Content-Disposition headers on every server.
    stem = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", stem).strip("-_")[:72]
    return f"{stem or 'transcript'}.{format}"


def _parts(meeting: dict) -> tuple[str, list[tuple[str, str]], list[dict]]:
    title = _clean(meeting.get("title")) or "Транскрипция встречи"
    metadata = [
        ("Идентификатор", _clean(meeting.get("id")) or "Не указан"),
        (
            "Источник",
            SOURCES.get(meeting.get("source"), _clean(meeting.get("source")) or "Не указан"),
        ),
        (
            "Язык",
            LANGUAGES.get(meeting.get("language"), _clean(meeting.get("language")) or "Не указан"),
        ),
        ("Создано", _date(meeting.get("created_at"))),
        ("Обновлено", _date(meeting.get("updated_at"))),
        ("Статус", _clean(meeting.get("status")) or "Не указан"),
    ]
    segments = meeting.get("segments") or []
    return title, metadata, segments


def _segment_heading(segment: dict) -> str:
    start, end = (
        _timestamp(segment.get("start")),
        _timestamp(segment.get("end", segment.get("start"))),
    )
    speaker = _clean(segment.get("speaker")) or "Спикер не указан"
    provisional = " · предварительный фрагмент" if not segment.get("final", True) else ""
    return f"[{start} — {end}] {speaker}{provisional}"


NOTICE = "Автоматическая транскрипция может содержать ошибки. Метки спикеров не подтверждают личность участника; проверьте текст перед использованием."


def _plain(meeting: dict) -> str:
    title, metadata, segments = _parts(meeting)
    lines = [title, "", *[f"{label}: {value}" for label, value in metadata], "", NOTICE, ""]
    if not segments:
        lines.append("Текст транскрипции пока отсутствует.")
    for segment in segments:
        lines.extend([_segment_heading(segment), _clean(segment.get("text")), ""])
    return "\n".join(lines).rstrip() + "\n"


def _md_escape(value: str) -> str:
    # Prevent transcript content from becoming active links, images, or HTML.
    value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    value = re.sub(r"([\\`*_{}\[\]()#|])", r"\\\1", value)
    return re.sub(r"(?m)^(\s*)([-+]|\d+\.) (?=\S)", r"\1\\\2 ", value)


def _markdown(meeting: dict) -> str:
    title, metadata, segments = _parts(meeting)
    lines = [
        f"# {_md_escape(title)}",
        "",
        *[f"- **{label}:** {_md_escape(value)}" for label, value in metadata],
        "",
        f"> {NOTICE}",
        "",
        "## Транскрипция",
        "",
    ]
    if not segments:
        lines.append("Текст транскрипции пока отсутствует.")
    for segment in segments:
        lines.extend(
            [
                f"### {_md_escape(_segment_heading(segment))}",
                "",
                _md_escape(_clean(segment.get("text"))),
                "",
            ]
        )
    return "\n".join(lines).rstrip() + "\n"


def _docx(meeting: dict) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    title, metadata, segments = _parts(meeting)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(8)
    document.core_properties.title = title
    document.core_properties.subject = "Транскрипция встречи"
    document.core_properties.author = "Meeting transcription platform"
    document.add_heading(title, 0)
    for label, value in metadata:
        p = document.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
    notice = document.add_paragraph(NOTICE)
    notice.runs[0].font.size = Pt(9)
    notice.runs[0].font.color.rgb = RGBColor.from_string("596579")
    document.add_heading("Транскрипция", 1)
    if not segments:
        document.add_paragraph("Текст транскрипции пока отсутствует.")
    for segment in segments:
        p = document.add_paragraph(_segment_heading(segment))
        p.paragraph_format.keep_with_next = True
        p.runs[0].bold = True
        document.add_paragraph(_clean(segment.get("text")))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _font_path() -> Path:
    custom = os.environ.get("EXPORT_FONT_PATH")
    candidates = [custom] if custom else []
    candidates += [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf",
        "/usr/local/share/fonts/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise RuntimeError(
        "Unicode PDF font missing. Install fonts-dejavu-core or set EXPORT_FONT_PATH to a TTF font supporting Kazakh, Russian and English."
    )


def _pdf(meeting: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    title, metadata, segments = _parts(meeting)
    path = _font_path()
    # Stable per-path name avoids collisions if an embedding app switches fonts.
    import hashlib

    font = "Transcript-" + hashlib.sha256(str(path).encode()).hexdigest()[:12]
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, str(path)))
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=(210 * mm, 297 * mm),
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=19 * mm,
        bottomMargin=20 * mm,
        title=title,
        author="Meeting transcription platform",
    )
    base = dict(
        fontName=font,
        leading=15,
        alignment=TA_LEFT,
        textColor=colors.HexColor("#192334"),
        splitLongWords=True,
    )
    styles = {
        "title": ParagraphStyle(
            "TranscriptTitle", fontSize=22, spaceAfter=18, **{**base, "leading": 28}
        ),
        "meta": ParagraphStyle("TranscriptMetadata", fontSize=9, spaceAfter=5, **base),
        "notice": ParagraphStyle(
            "TranscriptNotice",
            fontSize=8,
            spaceBefore=10,
            spaceAfter=18,
            **{**base, "leading": 12, "textColor": colors.HexColor("#606b7c")},
        ),
        "heading": ParagraphStyle(
            "TranscriptHeading",
            fontSize=10,
            spaceBefore=12,
            spaceAfter=5,
            keepWithNext=True,
            **base,
        ),
        "body": ParagraphStyle("TranscriptBody", fontSize=10, spaceAfter=6, **base),
    }

    def safe(value: str) -> str:
        return escape(value).replace("\n", "<br/>")

    elements = [Paragraph(safe(title), styles["title"])]
    elements += [Paragraph(safe(f"{label}: {value}"), styles["meta"]) for label, value in metadata]
    elements.append(Paragraph(safe(NOTICE), styles["notice"]))
    if not segments:
        elements.append(Paragraph("Текст транскрипции пока отсутствует.", styles["body"]))
    for segment in segments:
        elements.append(Paragraph(safe(_segment_heading(segment)), styles["heading"]))
        elements.append(Paragraph(safe(_clean(segment.get("text"))) or "&#160;", styles["body"]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#606b7c"))
        canvas.drawRightString(190 * mm, 11 * mm, str(doc.page))
        canvas.restoreState()

    document.build(elements, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_transcript(meeting: dict, format: str) -> tuple[bytes, str, str]:
    """Return download bytes, MIME type and an HTTP-header-safe filename.

    Only stored segments and metadata are exported. The caller owns authorization.
    """
    title = _clean(meeting.get("title")) or "transcript"
    format = str(format).lower()
    if format == "txt":
        content, mimetype = _plain(meeting).encode("utf-8"), "text/plain; charset=utf-8"
    elif format == "md":
        content, mimetype = _markdown(meeting).encode("utf-8"), "text/markdown; charset=utf-8"
    elif format == "docx":
        content, mimetype = (
            _docx(meeting),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    elif format == "pdf":
        content, mimetype = _pdf(meeting), "application/pdf"
    else:
        raise ValueError("Unsupported export format. Choose txt, md, docx or pdf.")
    return content, mimetype, _filename(title, format)
